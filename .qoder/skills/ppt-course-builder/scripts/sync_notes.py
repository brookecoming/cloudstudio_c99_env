"""把「第N章_口播稿.md」逐张回填进同章 PPT 的「演讲者备注」。

映射原理（唯一真理源 = 放映可见顺序）：
  python-pptx 的 prs.slides 顺序就是 presentation.xml 的 <p:sldIdLst> 放映序；
  跳过根元素带 show="0" 的隐藏页后，剩下的第 k 张 == 口播稿里的「🗂 第 k 张」。
  所以稿子怎么翻页、PPT 里怎么拖页序，重跑一次就自动对齐，不靠文件名、不靠编辑视图页码。

用法（pwsh 7 / Python 3）：
  python sync_notes.py "<口播稿.md>" "<课件.pptx>"            # dry-run：打印映射表 + KEEP/WRITE
  python sync_notes.py "<口播稿.md>" "<课件.pptx>" --apply     # 先备份再写入
  python sync_notes.py                                        # 不带参数：在**当前目录**自动找
                                                              #   第*_口播稿.md 与最大的 *.pptx
"""
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

from pptx import Presentation

# 口播稿的段落标题形如：### 🗂 第 12 张 · §2.4 算法的 5 个特性【约 3 分钟】
HEAD = re.compile(r"^###\s+🗂\s*第\s*(\d+)\s*张\s*·\s*(.*)$")
# 页头样式与既有备注保持一致，避免整本误判为“有改动”
NOTE_HEAD = "\u25b6 第 {n} 张 \u00b7 {title}\n{body}"


def clean(line: str) -> str:
    line = re.sub(r"^```.*$", "", line)      # 去掉代码围栏，保留代码正文（备注里要能看代码）
    return line.replace("**", "").rstrip()


def parse_script(md_text: str) -> dict[int, tuple[str, str]]:
    """返回 {放映可见页号: (该页标题, 正文)}；遇到 ## 一级节标题或 --- 分隔线即收尾该页。"""
    out: dict[int, tuple[str, str]] = {}
    num, title, buf = None, "", []

    def flush():
        if num is None:
            return
        body = "\n".join(clean(x) for x in buf).strip()
        out[num] = (title.strip(), re.sub(r"\n{3,}", "\n\n", body))

    for raw in md_text.splitlines():
        m = HEAD.match(raw)
        if m:
            flush()
            num, title, buf = int(m.group(1)), m.group(2), []
            continue
        if raw.startswith("## ") or raw.startswith("---"):
            flush()
            num, buf = None, []
            continue
        if num is None or raw.startswith(">"):   # blockquote=课间提示/尾注，不进单页备注
            continue
        buf.append(raw)
    flush()
    return out


def is_hidden(slide) -> bool:
    return slide._element.get("show") == "0"


def notes_text(slide) -> str:
    return slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""


def autodiscover() -> tuple[Path, Path] | None:
    """不带参数时：在当前目录找口播稿与课件（避开带空格路径在 shell 里被拆坏的坑）。"""
    here = Path.cwd()
    mds = sorted(here.glob("*口播稿*.md")) or sorted(here.glob("*提词本*.md"))
    pptxs = [p for p in here.glob("*.pptx") if not p.name.startswith("~$")]
    if not mds or not pptxs:
        return None
    pptx = max(pptxs, key=lambda p: p.stat().st_size)      # 取最大的，避开分卷/占位文件
    return mds[0], pptx


def main(argv: list[str]) -> int:
    apply = "--apply" in argv
    args = [a for a in argv if not a.startswith("--")]
    if len(args) == 2:
        md, pptx = Path(args[0]), Path(args[1])
    elif not args:
        found = autodiscover()
        if not found:
            print(__doc__)
            print("!! 当前目录没找到『*口播稿*.md + *.pptx』，请显式传两个路径")
            return 2
        md, pptx = found
        print(f"[自动发现] {md.name}  →  {pptx.name}")
    else:
        print(__doc__)
        return 2
    for p in (md, pptx):
        if not p.exists():
            print(f"!! 找不到文件：{p}")
            return 2

    script = parse_script(md.read_text(encoding="utf-8"))
    nums = sorted(script)
    if nums != list(range(1, len(nums) + 1)):
        print(f"!! 稿里页号不连续（{nums}），先修稿再回填")
        return 1

    prs = Presentation(str(pptx))
    visible = [(i, s) for i, s in enumerate(prs.slides, start=1) if not is_hidden(s)]
    print(f"[稿 ] {md.name} 解析出 {len(nums)} 段，页号 1..{nums[-1]}")
    print(f"[PPT] {pptx.name} 绝对 {len(prs.slides._sldIdLst)} 张 / 隐藏 {len(prs.slides._sldIdLst) - len(visible)} 张 / 可见 {len(visible)} 张")
    if len(visible) != len(nums):
        print("!! 可见页数与稿段数不一致：说明稿还没跟上课件结构，先按课件改稿（提词本永远最后写），别硬写")
        return 1

    todo = []
    for k, (abs_idx, slide) in enumerate(visible, start=1):
        title, body = script[k]
        body = NOTE_HEAD.format(n=k, title=title, body=body)
        script[k] = (title, body)
        mark = "KEEP" if notes_text(slide).strip() == body.strip() else "WRITE"
        if mark == "WRITE":
            todo.append(k)
        print(f"{k:>2} → abs{abs_idx:<3} {mark:<5} {title[:30]}")
    print(f"\n需写入 {len(todo)} 页：{todo}")

    if not apply:
        print("dry-run 完成，未改动文件；确认后加 --apply。")
        return 0

    bak = pptx.with_name(pptx.name + ".bak-" + datetime.now().strftime("%Y%m%d_%H%M%S"))
    shutil.copy2(pptx, bak)
    print(f"[备份] {bak.name}")
    for k, (abs_idx, slide) in enumerate(visible, start=1):
        body = script[k][1]
        if notes_text(slide).strip() == body.strip():
            continue
        slide.notes_slide.notes_text_frame.text = body
    prs.save(str(pptx))
    print(f"[写入] 已保存 {len(todo)} 页备注。请重跑一次 dry-run，应显示『需写入 0 页』。")
    return 0


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")  # Windows 控制台默认 GBK，中文/箭头符号会炸
    sys.exit(main(sys.argv[1:]))
